"""Parse Xiaoyu's XLSX replies for the 07_待確認事項 sheet."""

from __future__ import annotations

import argparse
import logging
import os
import re
import sqlite3
import sys
from typing import Any

import openpyxl

ROOT_DIR = os.path.dirname(os.path.dirname(__file__))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

DB_PATH = os.path.join(ROOT_DIR, "data", "shanbot.db")
SHEET_NAME = "07_待確認事項"
CANDIDATE_RE = re.compile(r"^(.+?)（[\d.]+分）$")

logger = logging.getLogger("shanbot.email_reply_parser")


def _get_conn() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _company_id(value: Any) -> int | None:
    text = _cell_text(value)
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _parse_candidate(value: Any) -> str:
    text = _cell_text(value)
    if not text:
        return ""
    match = CANDIDATE_RE.match(text)
    return match.group(1).strip() if match else ""


def _find_remark_col(ws) -> int:
    for header_row in range(1, min(ws.max_row, 5) + 1):
        for col in range(15, ws.max_column + 1):
            header = _cell_text(ws.cell(header_row, col).value)
            if any(key in header for key in ("備註", "正式名稱", "新供應商")):
                return col
    return 15


def _clean_remark_name(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^(新供應商|正式名稱|供應商名稱|名稱|備註)\s*[:：]\s*", "", text)
    return text.strip()


def _extract_new_supplier_name(choice: str, remark: str) -> str:
    if remark:
        return _clean_remark_name(remark)

    tail = re.sub(r"^都不是\s*", "", choice).strip()
    tail = tail.lstrip("：:，,。.- ")
    return _clean_remark_name(tail)


def get_or_create_supplier(name: str) -> dict:
    """Return {'id', 'name', 'created'} for a supplier canonical name."""
    conn = _get_conn()
    try:
        row = conn.execute(
            "SELECT id, name FROM suppliers WHERE name = ?",
            (name,),
        ).fetchone()
        if row:
            return {"id": row["id"], "name": row["name"], "created": False}

        cur = conn.execute(
            "INSERT INTO suppliers (name, has_uniform_invoice) VALUES (?, 0)",
            (name,),
        )
        conn.commit()
        return {"id": cur.lastrowid, "name": name, "created": True}
    finally:
        conn.close()


def supplier_exists(name: str) -> bool:
    conn = _get_conn()
    try:
        row = conn.execute("SELECT 1 FROM suppliers WHERE name = ?", (name,)).fetchone()
        return row is not None
    finally:
        conn.close()


def update_staging_by_name(
    company_id: int,
    supplier_name: str,
    supplier_id: int,
    canonical_name: str,
) -> int:
    """Batch-update unresolved staging rows by company and OCR supplier name."""
    conn = _get_conn()
    try:
        cur = conn.execute(
            """
            UPDATE purchase_staging
            SET supplier_id = ?, supplier_name = ?
            WHERE company_id = ?
              AND supplier_name = ?
              AND supplier_id IS NULL
            """,
            (supplier_id, canonical_name, company_id, supplier_name),
        )
        conn.commit()
        return cur.rowcount
    finally:
        conn.close()


def count_staging_by_name(company_id: int, supplier_name: str) -> int:
    conn = _get_conn()
    try:
        row = conn.execute(
            """
            SELECT COUNT(*) AS cnt
            FROM purchase_staging
            WHERE company_id = ?
              AND supplier_name = ?
              AND supplier_id IS NULL
            """,
            (company_id, supplier_name),
        ).fetchone()
        return int(row["cnt"] if row else 0)
    finally:
        conn.close()


def write_alias(alias_text: str, supplier_id: int, week: str, employee_id: str) -> bool:
    """Write supplier alias with confidence=100 and upsert on alias_text."""
    if not alias_text or not supplier_id:
        return False

    conn = _get_conn()
    try:
        conn.execute(
            """
            INSERT INTO suppliers_alias
                (alias_text, canonical_supplier_id, confidence,
                 learned_from, learned_at, employee_id)
            VALUES (?, ?, 100, ?, datetime('now','localtime'), ?)
            ON CONFLICT(alias_text) DO UPDATE SET
                canonical_supplier_id = excluded.canonical_supplier_id,
                confidence = 100,
                learned_from = excluded.learned_from,
                learned_at = excluded.learned_at,
                employee_id = excluded.employee_id
            """,
            (alias_text, supplier_id, f"email_reply:{week or ''}", employee_id or ""),
        )
        conn.commit()
        return True
    finally:
        conn.close()


def _resolve_choice(choice: str, candidates: list[str], remark: str) -> tuple[str, str]:
    if choice in {"1", "2", "3"}:
        idx = int(choice) - 1
        if idx >= len(candidates) or not candidates[idx]:
            return "", ""
        return candidates[idx], "candidate"

    if choice.startswith("都不是"):
        new_name = _extract_new_supplier_name(choice, remark)
        return new_name, "new_from_remark"

    return choice, "formal_name"


def parse_xlsx_reply(
    xlsx_path: str,
    week: str = "",
    employee_id: str = "",
    thread_id: str = "",
    dry_run: bool = False,
    create_reply_draft: bool = True,
) -> dict:
    """Parse and apply supplier alias replies from 07_待確認事項."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Workbook missing sheet: {SHEET_NAME}")

    ws = wb[SHEET_NAME]
    remark_col = _find_remark_col(ws)

    aliases_learned = 0
    staging_updates = 0
    new_suppliers = 0
    skipped = 0
    rows_processed = 0
    processed_rows: list[dict] = []

    for row_idx in range(4, ws.max_row + 1):
        company_id = _company_id(ws.cell(row_idx, 1).value)
        supplier_name = _cell_text(ws.cell(row_idx, 4).value)
        choice = _cell_text(ws.cell(row_idx, 14).value)

        if not any((company_id, supplier_name, choice)):
            continue
        if not company_id or not supplier_name or not choice:
            skipped += 1
            continue

        candidates = [_parse_candidate(ws.cell(row_idx, col).value) for col in (11, 12, 13)]
        remark = _cell_text(ws.cell(row_idx, remark_col).value)
        canonical_name, choice_type = _resolve_choice(choice, candidates, remark)

        if not canonical_name:
            skipped += 1
            continue

        rows_processed += 1

        if dry_run:
            supplier_id = -1
            created = not supplier_exists(canonical_name)
            updated_count = count_staging_by_name(company_id, supplier_name)
            alias_ok = True
        else:
            supplier = get_or_create_supplier(canonical_name)
            supplier_id = supplier["id"]
            created = supplier["created"]
            updated_count = update_staging_by_name(
                company_id=company_id,
                supplier_name=supplier_name,
                supplier_id=supplier_id,
                canonical_name=canonical_name,
            )
            alias_ok = write_alias(
                alias_text=supplier_name,
                supplier_id=supplier_id,
                week=week,
                employee_id=employee_id,
            )

        if alias_ok:
            aliases_learned += 1
        staging_updates += updated_count
        if created:
            new_suppliers += 1

        processed_rows.append(
            {
                "row": row_idx,
                "company_id": company_id,
                "supplier_name": supplier_name,
                "supplier_id": supplier_id,
                "canonical": canonical_name,
                "choice": choice,
                "choice_type": choice_type,
                "remark": remark,
                "staging_updates": updated_count,
            }
        )

    draft_id = None
    if create_reply_draft and not dry_run:
        from handlers.email_reply_draft import create_reply_draft as _create_reply_draft

        draft_id = _create_reply_draft(
            week=week,
            thread_id=thread_id,
            aliases_learned=aliases_learned,
            staging_updates=staging_updates,
            new_suppliers=new_suppliers,
            processed_rows=processed_rows,
        )

    return {
        "aliases_learned": aliases_learned,
        "staging_updates": staging_updates,
        "new_suppliers": new_suppliers,
        "skipped": skipped,
        "rows_processed": rows_processed,
        "draft_id": draft_id,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Parse Xiaoyu XLSX supplier replies")
    parser.add_argument("xlsx_path", nargs="?", help="XLSX reply file path")
    parser.add_argument("--week", default="", help="Week label, e.g. W20")
    parser.add_argument("--employee-id", default="", help="Employee id for alias learning")
    parser.add_argument("--thread-id", default="", help="Gmail thread id for reply draft")
    parser.add_argument("--dry-run", action="store_true", help="Parse without DB writes or draft")
    parser.add_argument("--no-draft", action="store_true", help="Do not create Gmail reply draft")
    parser.add_argument("--mock", action="store_true", help="Run tests.test_email_reply_parser.run_mock_test")
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")

    if args.mock:
        from tests.test_email_reply_parser import run_mock_test

        run_mock_test()
        return

    if not args.xlsx_path:
        parser.error("xlsx_path is required unless --mock is used")

    result = parse_xlsx_reply(
        xlsx_path=args.xlsx_path,
        week=args.week,
        employee_id=args.employee_id,
        thread_id=args.thread_id,
        dry_run=args.dry_run,
        create_reply_draft=not args.no_draft,
    )
    print(result)


if __name__ == "__main__":
    main()
