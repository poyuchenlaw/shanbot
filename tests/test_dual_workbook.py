"""雙版本帳冊測試（員工驗章 / 小魚決策）

驗收條件：
1. 員工版 4 sheets，索引在第 0 位，禁字串「試算/分錄/借方/貸方/損益/資負/總帳」
2. 老闆版 ≥7 sheets，索引在第 0 位
3. 兩份檔名共存同 dir 互不覆寫
4. variant 參數驗證
"""

from __future__ import annotations

import importlib
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook

BANNED_STAFF_WORDS = [
    "試算", "分錄", "借方", "貸方", "損益", "資負", "總帳",
    "journal", "ledger", "trial_balance",
]


def _seed_test_data(tmpdir: str):
    """初始化測試 DB + 注入 3 筆 staging (confirmed/pending/supplier空白)"""
    db_path = os.path.join(tmpdir, "test_shanbot.db")
    os.environ["DB_PATH"] = db_path

    import state_manager
    importlib.reload(state_manager)
    state_manager.init_db()

    sid1 = state_manager.add_purchase_staging(
        user_id="U_test", chat_id="U_test",
        purchase_date="2026-04-05", company_id=1,
    )
    state_manager.update_purchase_staging(
        sid1, supplier_name="阿榮蔬果", total_amount=3500, tax_amount=175,
    )
    state_manager.confirm_staging(sid1)

    sid2 = state_manager.add_purchase_staging(
        user_id="U_test", chat_id="U_test",
        purchase_date="2026-04-12", company_id=1,
    )
    state_manager.update_purchase_staging(
        sid2, supplier_name="義美食品", total_amount=1200, tax_amount=60,
    )

    sid3 = state_manager.add_purchase_staging(
        user_id="U_test", chat_id="U_test",
        purchase_date="2026-04-15", company_id=1,
    )
    state_manager.update_purchase_staging(
        sid3, supplier_name="", total_amount=0, tax_amount=0,
    )

    # 重新 reload services（讓 ACCOUNTING_DIR 之外的 module 拿到新 DB_PATH）
    import services.accounting_service
    importlib.reload(services.accounting_service)
    import services.staff_workbook_service
    importlib.reload(services.staff_workbook_service)


def _patch_dir(tmpdir):
    import services.accounting_service as acct
    orig = acct.ACCOUNTING_DIR
    acct.ACCOUNTING_DIR = tmpdir
    return orig


def _restore_dir(orig):
    import services.accounting_service as acct
    acct.ACCOUNTING_DIR = orig


class TestDualWorkbookStaff(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        _seed_test_data(self.tmpdir)
        self._orig_dir = _patch_dir(self.tmpdir)

    def tearDown(self):
        _restore_dir(self._orig_dir)

    def test_staff_workbook_4_sheets(self):
        from services.accounting_service import generate_accounting_excel
        path = generate_accounting_excel("2026-04", variant="staff")
        self.assertTrue(path and os.path.exists(path))
        wb = load_workbook(path)
        sheets = wb.sheetnames
        self.assertEqual(len(sheets), 4,
                         f"員工版必須 4 sheets, 實際 {len(sheets)}: {sheets}")
        self.assertIn("索引", sheets[0],
                       f"第 0 個 sheet 必須是索引, 實際: {sheets[0]}")

    def test_staff_filename_correct(self):
        from services.accounting_service import generate_accounting_excel
        path = generate_accounting_excel("2026-04", variant="staff")
        self.assertIn("員工驗章帳冊", os.path.basename(path))

    def test_staff_no_accounting_jargon(self):
        from services.accounting_service import generate_accounting_excel
        path = generate_accounting_excel("2026-04", variant="staff")
        wb = load_workbook(path)

        for sheet_name in wb.sheetnames:
            for banned in BANNED_STAFF_WORDS:
                self.assertNotIn(
                    banned, sheet_name,
                    f"sheet title '{sheet_name}' 含禁字串 '{banned}'")

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows(values_only=True):
                for cell_value in row:
                    if cell_value is None:
                        continue
                    text = str(cell_value)
                    for banned in BANNED_STAFF_WORDS:
                        self.assertNotIn(
                            banned, text,
                            f"sheet '{sheet_name}' cell 含禁字串 '{banned}': {text[:80]}")

    def test_staff_includes_pending(self):
        from services.accounting_service import generate_accounting_excel
        path = generate_accounting_excel("2026-04", variant="staff")
        wb = load_workbook(path)
        verify_sheet = None
        for name in wb.sheetnames:
            if "驗章" in name:
                verify_sheet = wb[name]
                break
        self.assertIsNotNone(verify_sheet, "找不到驗章清單 sheet")

        all_text = []
        for row in verify_sheet.iter_rows(values_only=True):
            for c in row:
                if c is not None:
                    all_text.append(str(c))
        joined = " ".join(all_text)
        self.assertIn("V", joined, "驗章清單必含 V 標記")
        self.assertIn("X", joined, "驗章清單必含 X 標記")


class TestDualWorkbookBoss(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        _seed_test_data(self.tmpdir)
        self._orig_dir = _patch_dir(self.tmpdir)

    def tearDown(self):
        _restore_dir(self._orig_dir)

    def test_boss_workbook_has_index_at_pos_0(self):
        from services.accounting_service import generate_accounting_excel
        path = generate_accounting_excel("2026-04", variant="boss")
        self.assertTrue(path and os.path.exists(path))
        wb = load_workbook(path)
        sheets = wb.sheetnames
        self.assertGreaterEqual(len(sheets), 7,
                                  f"老闆版至少 7 sheets, 實際 {len(sheets)}: {sheets}")
        self.assertIn("索引", sheets[0],
                       f"第 0 個 sheet 必須是索引, 實際: {sheets[0]}")

    def test_boss_filename_correct(self):
        from services.accounting_service import generate_accounting_excel
        path = generate_accounting_excel("2026-04", variant="boss")
        self.assertIn("小魚決策帳冊", os.path.basename(path))


class TestDualWorkbookCoexist(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        _seed_test_data(self.tmpdir)
        self._orig_dir = _patch_dir(self.tmpdir)

    def tearDown(self):
        _restore_dir(self._orig_dir)

    def test_both_variants_coexist_same_dir(self):
        from services.accounting_service import generate_accounting_excel
        boss_path = generate_accounting_excel("2026-04", variant="boss")
        staff_path = generate_accounting_excel("2026-04", variant="staff")

        self.assertTrue(os.path.exists(boss_path), "老闆版檔案不存在")
        self.assertTrue(os.path.exists(staff_path), "員工版檔案不存在")
        self.assertNotEqual(boss_path, staff_path, "兩份檔名相同（衝突）")
        self.assertEqual(os.path.dirname(boss_path),
                         os.path.dirname(staff_path),
                         "兩份不在同資料夾")

    def test_invalid_variant_raises(self):
        from services.accounting_service import generate_accounting_excel
        with self.assertRaises(ValueError):
            generate_accounting_excel("2026-04", variant="invalid")


if __name__ == "__main__":
    unittest.main(verbosity=2)
