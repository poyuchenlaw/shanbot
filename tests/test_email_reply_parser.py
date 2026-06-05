import os
import sys
import logging
import sqlite3
import tempfile

import openpyxl
import pytest

ROOT = os.path.dirname(os.path.dirname(__file__))
sys.path.insert(0, ROOT)

DB_PATH = os.path.join(ROOT, "data", "shanbot.db")

from handlers.email_reply_parser import parse_xlsx_reply, _parse_candidate, _resolve_choice


logger = logging.getLogger(__name__)


def _build_mock_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "07_待確認事項"

    headers = {
        1: "公司代碼",
        4: "供應商名稱",
        11: "候選1",
        12: "候選2",
        13: "候選3",
        14: "妳的選擇",
    }
    for col in range(1, 15):
        ws.cell(3, col).value = headers.get(col, f"欄位{col}")

    rows = [
        (4, 1, "圻逸食品", "圻逸食品有限公司（100.0分）", "忻逸食品（87.5分）", "昕逸食品（80.0分）", "1"),
        (5, 1, "黃子儀蔬菜", "黃子儀蔬菜物流（100.0分）", "興丁蔬菜（61.5分）", "雙子星（16.6分）", "2"),
        (6, 1, "詠源順環保", "詠源順環保有限公司（100.0分）", "潮崴有限公司（53.3分）", "瀚崴有限公司（53.3分）", "3"),
        (7, 1, "晨昌", "晨昌公司（100.0分）", "晨星公司（70.0分）", "晨光公司（60.0分）", "晨昌公司"),
        (8, 1, "不明廠商", "候選A（50.0分）", "候選B（40.0分）", "候選C（30.0分）", "都不是:新大陸食品行"),
    ]
    for row_idx, company_id, supplier_name, c11, c12, c13, choice in rows:
        ws.cell(row_idx, 1).value = company_id
        ws.cell(row_idx, 4).value = supplier_name
        ws.cell(row_idx, 11).value = c11
        ws.cell(row_idx, 12).value = c12
        ws.cell(row_idx, 13).value = c13
        ws.cell(row_idx, 14).value = choice

    wb.save(path)


def _email_reply_alias_count():
    conn = sqlite3.connect(DB_PATH)
    try:
        row = conn.execute(
            "SELECT COUNT(*) FROM suppliers_alias WHERE learned_from LIKE ?",
            ("email_reply%",),
        ).fetchone()
        return int(row[0])
    finally:
        conn.close()


def run_mock_test(dry_run=False):
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    try:
        _build_mock_xlsx(tmp.name)
        result = parse_xlsx_reply(
            tmp.name,
            week="W20",
            employee_id="uiy022803",
            thread_id="",
            dry_run=dry_run,
            create_reply_draft=False,
        )

        assert result["rows_processed"] >= 5
        if not dry_run:
            assert result["aliases_learned"] >= 1
            alias_count = _email_reply_alias_count()
            assert alias_count >= result["aliases_learned"]
        else:
            alias_count = None

        print(
            "email_reply_parser mock: "
            f"dry_run={dry_run} rows_processed={result['rows_processed']} "
            f"aliases_learned={result['aliases_learned']} "
            f"email_reply_alias_count={alias_count}"
        )
        return result
    finally:
        os.unlink(tmp.name)


def test_parse_candidate():
    assert _parse_candidate("圻逸食品有限公司（100.0分）") == "圻逸食品有限公司"
    assert _parse_candidate(" 黃子儀蔬菜物流（87.5分） ") == "黃子儀蔬菜物流"
    assert _parse_candidate("候選A(50.0分)") == ""
    assert _parse_candidate("候選A（分）") == ""
    assert _parse_candidate("") == ""
    assert _parse_candidate(None) == ""


def test_resolve_choice_variants():
    candidates = ["候選A", "候選B", "候選C"]

    assert _resolve_choice("1", candidates, "") == ("候選A", "candidate")
    assert _resolve_choice("2", candidates, "") == ("候選B", "candidate")
    assert _resolve_choice("3", candidates, "") == ("候選C", "candidate")
    assert _resolve_choice("晨昌公司", candidates, "") == ("晨昌公司", "formal_name")
    assert _resolve_choice("都不是:新大陸食品行", candidates, "") == ("新大陸食品行", "new_from_remark")


def test_mock_dry_run():
    result = run_mock_test(dry_run=True)
    assert result["rows_processed"] >= 5


def test_mock_write_db():
    result = run_mock_test(dry_run=False)
    assert result["aliases_learned"] >= 1
    assert _email_reply_alias_count() >= result["aliases_learned"]


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    print(run_mock_test(dry_run=False))
