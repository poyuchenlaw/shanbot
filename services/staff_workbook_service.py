"""員工驗章帳冊（staff verification workbook）— 4 sheets

Sheets:
    📋 索引 / ✅ 驗章清單 / 🏢 本月供應商彙總 / ⚠️ 異常待釐清

資料源：confirmed (V) + pending (X) 都顯示
"""

from __future__ import annotations

import logging
import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import state_manager as sm

logger = logging.getLogger("shanbot.staff_workbook")

_HEADER_FONT = Font(name="微軟正黑體", size=11, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
_TITLE_FONT = Font(name="微軟正黑體", size=14, bold=True)
_NORMAL_FONT = Font(name="微軟正黑體", size=10)
_MONEY_FMT = "#,##0"
_THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
_CENTER = Alignment(horizontal="center", vertical="center")
_RIGHT = Alignment(horizontal="right", vertical="center")
_LIGHT_GREEN = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
_LIGHT_RED = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")

CONFIRMED_STATES = ("confirmed", "reported", "exported", "archived")
STATUS_LABEL_MAP = {
    "pending": "待處理",
    "confirmed": "已入帳",
    "reported": "已入帳",
    "exported": "已入帳",
    "archived": "已入帳",
    "discarded": "已捨棄",
}


def _ensure_dir(path: str):
    os.makedirs(path, exist_ok=True)


def generate_staff_workbook(year_month: str, company_id: int = None,
                            output_dir: str = None) -> str:
    """員工驗章帳冊（4 sheets）

    Args:
        year_month: YYYY-MM
        company_id: 指定公司（None = 全公司合併）
        output_dir: 輸出目錄（None = 使用 ACCOUNTING_DIR/{ym}/）

    Returns: Excel 檔案路徑
    """
    if output_dir is None:
        from services.accounting_service import ACCOUNTING_DIR
        output_dir = os.path.join(ACCOUNTING_DIR, year_month)
    _ensure_dir(output_dir)
    company_suffix = f"_C{company_id}" if company_id else ""
    filepath = os.path.join(
        output_dir, f"{year_month}_員工驗章帳冊{company_suffix}.xlsx")

    confirmed = sm.get_stagings_by_month(year_month, company_id=company_id)
    pending_all = sm.get_pending_stagings(company_id=company_id)
    pending = [s for s in pending_all
               if (s.get("year_month") or "").startswith(year_month)]
    all_stagings = confirmed + pending

    wb = openpyxl.Workbook()
    ws_idx = wb.active
    ws_idx.title = "📋 索引"
    _build_index_sheet(
        ws_idx, year_month=year_month, company_id=company_id,
        sheet_specs=[
            ("✅ 驗章清單",
             "本月所有收據逐筆列出，已入帳=V / 待處理=X",
             "你每天上傳完後在這裡勾稽，確認沒漏單"),
            ("🏢 本月供應商彙總",
             "按供應商彙總筆數 / 金額（含已入帳與待處理）",
             "一頁看完本月跟哪些廠商往來、各花多少"),
            ("⚠️ 異常待釐清",
             "供應商空白 / 金額為 0 / 超過 14 天未處理",
             "看到這頁有筆數 → 請補上資料或回報小魚"),
        ],
    )

    _write_verification_list(wb.create_sheet("✅ 驗章清單"),
                              all_stagings, year_month)
    _write_supplier_summary(wb.create_sheet("🏢 本月供應商彙總"),
                             all_stagings, year_month)
    _write_anomalies(wb.create_sheet("⚠️ 異常待釐清"),
                      all_stagings, year_month)

    try:
        from services.excel_merge import save_with_shadow
        save_with_shadow(wb, filepath)
    except ImportError:
        wb.save(filepath)
    logger.info(f"Staff verification workbook generated: {filepath}")
    return filepath


def _build_index_sheet(ws, year_month: str, company_id, sheet_specs: list):
    """套用 5/17 Excel 索引憲法：前 3-5 行說明區 + 分頁列表 + 對比價值"""
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 42
    ws.column_dimensions["D"].width = 38

    ws.merge_cells("A1:D1")
    ws["A1"] = f"員工驗章帳冊 — {year_month}"
    ws["A1"].font = Font(name="微軟正黑體", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(start_color="1565C0", end_color="1565C0", fill_type="solid")
    ws["A1"].alignment = _CENTER
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:D2")
    ws["A2"] = "用途：員工每天勾稽 / 月底交給小魚前的自查"
    ws["A2"].font = Font(name="微軟正黑體", size=11, italic=True)
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.merge_cells("A3:D3")
    ws["A3"] = ("你只要做兩件事：(1) 在『✅ 驗章清單』勾 V/X　"
                "(2) 看『⚠️ 異常待釐清』有沒有筆數")
    ws["A3"].font = Font(name="微軟正黑體", size=11, bold=True, color="C62828")
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 22

    company_label = f"公司 ID={company_id}" if company_id else "全公司合併"
    gen_at = datetime.now().strftime("%Y-%m-%d %H:%M")
    ws.merge_cells("A4:D4")
    ws["A4"] = f"範圍：{company_label} ／ 期間：{year_month} ／ 產出時間：{gen_at}"
    ws["A4"].font = Font(name="微軟正黑體", size=10, color="666666")
    ws["A4"].alignment = Alignment(horizontal="left")

    headers = ["#", "分頁名稱", "分頁意義（這頁告訴你什麼）",
               "對比價值（為什麼這頁不可少）"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col_idx, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _THIN_BORDER

    for i, (name, meaning, value) in enumerate(sheet_specs, start=1):
        r = 6 + i
        ws.cell(row=r, column=1, value=i).font = _NORMAL_FONT
        ws.cell(row=r, column=2, value=name).font = Font(name="微軟正黑體", size=10, bold=True)
        ws.cell(row=r, column=3, value=meaning).font = _NORMAL_FONT
        ws.cell(row=r, column=4, value=value).font = _NORMAL_FONT
        ws.row_dimensions[r].height = 30
        for col in range(1, 5):
            c = ws.cell(row=r, column=col)
            c.border = _THIN_BORDER
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _write_verification_list(ws, stagings: list, year_month: str):
    """Sheet 2: ✅ 驗章清單"""
    ws.merge_cells("A1:H1")
    ws["A1"] = f"✅ 驗章清單 — {year_month}（請逐筆檢查）"
    ws["A1"].font = _TITLE_FONT
    ws["A1"].alignment = _CENTER

    ws.merge_cells("A2:H2")
    ws["A2"] = "勾稽方式：已入帳=V ／ 待處理=X ／ 看到 X 請聯絡小魚或重新拍照"
    ws["A2"].font = Font(name="微軟正黑體", size=10, italic=True, color="C62828")

    headers = ["收據編號", "採購日期", "供應商", "金額", "狀態", "上傳時間", "備註", "勾稽"]
    widths = [10, 12, 22, 12, 10, 18, 24, 8]
    for col_idx, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    row = 5
    for s in sorted(stagings, key=lambda x: (x.get("purchase_date") or "", x.get("id", 0))):
        status = s.get("status") or "pending"
        mark = "V" if status in CONFIRMED_STATES else "X"
        status_label = STATUS_LABEL_MAP.get(status, status)

        ws.cell(row=row, column=1, value=f"#{s.get('id', '')}").font = _NORMAL_FONT
        ws.cell(row=row, column=2, value=s.get("purchase_date") or "").font = _NORMAL_FONT
        ws.cell(row=row, column=3, value=s.get("supplier_name") or "（空白）").font = _NORMAL_FONT
        amt_cell = ws.cell(row=row, column=4, value=s.get("total_amount") or 0)
        amt_cell.font = _NORMAL_FONT
        amt_cell.number_format = _MONEY_FMT
        amt_cell.alignment = _RIGHT
        status_cell = ws.cell(row=row, column=5, value=status_label)
        status_cell.font = _NORMAL_FONT
        status_cell.fill = _LIGHT_GREEN if mark == "V" else _LIGHT_RED
        ws.cell(row=row, column=6, value=s.get("created_at") or "").font = _NORMAL_FONT
        ws.cell(row=row, column=7, value=s.get("notes") or "").font = _NORMAL_FONT
        mark_cell = ws.cell(row=row, column=8, value=mark)
        mark_cell.font = Font(name="微軟正黑體", size=12, bold=True,
                              color="2E7D32" if mark == "V" else "C62828")
        mark_cell.alignment = _CENTER
        for c in range(1, 9):
            ws.cell(row=row, column=c).border = _THIN_BORDER
        row += 1

    if row == 5:
        ws.merge_cells(f"A{row}:H{row}")
        ws.cell(row=row, column=1, value="（本月無收據資料）").font = _NORMAL_FONT


def _write_supplier_summary(ws, stagings: list, year_month: str):
    """Sheet 3: 🏢 本月供應商彙總"""
    ws.merge_cells("A1:E1")
    ws["A1"] = f"🏢 本月供應商彙總 — {year_month}"
    ws["A1"].font = _TITLE_FONT
    ws["A1"].alignment = _CENTER

    ws.merge_cells("A2:E2")
    ws["A2"] = ("用途：一頁看完本月跟哪些廠商往來、各花多少。"
                "資料含『已入帳』與『待處理』兩種狀態。")
    ws["A2"].font = Font(name="微軟正黑體", size=10, italic=True)

    headers = ["供應商", "已入帳筆數", "待處理筆數", "本月金額合計", "備註"]
    widths = [26, 12, 12, 16, 24]
    for col_idx, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    summary: dict = {}
    for s in stagings:
        name = s.get("supplier_name") or "（空白）"
        if name not in summary:
            summary[name] = {"v": 0, "x": 0, "amount": 0}
        status = s.get("status") or "pending"
        if status in CONFIRMED_STATES:
            summary[name]["v"] += 1
        elif status == "pending":
            summary[name]["x"] += 1
        summary[name]["amount"] += s.get("total_amount") or 0

    row = 5
    for name, data in sorted(summary.items(), key=lambda x: -x[1]["amount"]):
        ws.cell(row=row, column=1, value=name).font = _NORMAL_FONT
        ws.cell(row=row, column=2, value=data["v"]).font = _NORMAL_FONT
        x_cell = ws.cell(row=row, column=3, value=data["x"])
        x_cell.font = _NORMAL_FONT
        if data["x"] > 0:
            x_cell.fill = _LIGHT_RED
        amt_cell = ws.cell(row=row, column=4, value=data["amount"])
        amt_cell.font = _NORMAL_FONT
        amt_cell.number_format = _MONEY_FMT
        amt_cell.alignment = _RIGHT
        note = "有 X 待處理" if data["x"] > 0 else "全部完成"
        ws.cell(row=row, column=5, value=note).font = _NORMAL_FONT
        for c in range(1, 6):
            ws.cell(row=row, column=c).border = _THIN_BORDER
        row += 1

    if row == 5:
        ws.merge_cells(f"A{row}:E{row}")
        ws.cell(row=row, column=1, value="（本月無供應商資料）").font = _NORMAL_FONT


def _write_anomalies(ws, stagings: list, year_month: str):
    """Sheet 4: ⚠️ 異常待釐清"""
    ws.merge_cells("A1:F1")
    ws["A1"] = f"⚠️ 異常待釐清 — {year_month}"
    ws["A1"].font = _TITLE_FONT
    ws["A1"].alignment = _CENTER

    ws.merge_cells("A2:F2")
    ws["A2"] = "看到這頁有筆數 → 請補上資料或回報小魚。沒筆數 = 本月乾淨。"
    ws["A2"].font = Font(name="微軟正黑體", size=10, italic=True, color="C62828")

    headers = ["收據編號", "採購日期", "供應商", "金額", "異常類型", "建議動作"]
    widths = [10, 12, 22, 12, 18, 30]
    for col_idx, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    today = datetime.now()
    row = 5
    found = False

    for s in stagings:
        issues = []
        actions = []
        supplier = s.get("supplier_name") or ""
        amount = s.get("total_amount") or 0
        status = s.get("status") or "pending"
        created = s.get("created_at") or ""

        if not supplier or supplier in ("未知", "（空白）"):
            issues.append("供應商空白")
            actions.append("補上供應商名稱")
        if amount <= 0:
            issues.append("金額為 0")
            actions.append("確認收據金額")
        if status == "pending":
            try:
                created_dt = datetime.strptime(created[:10], "%Y-%m-%d")
                days = (today - created_dt).days
                if days >= 14:
                    issues.append(f"待處理 {days} 天")
                    actions.append("請小魚或老闆儘速確認")
            except (ValueError, TypeError):
                pass

        if not issues:
            continue
        found = True

        ws.cell(row=row, column=1, value=f"#{s.get('id', '')}").font = _NORMAL_FONT
        ws.cell(row=row, column=2, value=s.get("purchase_date") or "").font = _NORMAL_FONT
        ws.cell(row=row, column=3, value=supplier or "（空白）").font = _NORMAL_FONT
        amt_cell = ws.cell(row=row, column=4, value=amount)
        amt_cell.font = _NORMAL_FONT
        amt_cell.number_format = _MONEY_FMT
        amt_cell.alignment = _RIGHT
        issue_cell = ws.cell(row=row, column=5, value=" / ".join(issues))
        issue_cell.font = _NORMAL_FONT
        issue_cell.fill = _LIGHT_RED
        ws.cell(row=row, column=6, value=" / ".join(actions)).font = _NORMAL_FONT
        for c in range(1, 7):
            ws.cell(row=row, column=c).border = _THIN_BORDER
        row += 1

    if not found:
        ws.merge_cells(f"A{row}:F{row}")
        cell = ws.cell(row=row, column=1, value="✅ 本月無異常，全部資料完整")
        cell.font = Font(name="微軟正黑體", size=12, bold=True, color="2E7D32")
        cell.alignment = _CENTER
